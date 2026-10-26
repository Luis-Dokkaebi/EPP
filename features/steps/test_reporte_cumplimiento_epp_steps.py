"""Pasos de features/reporte_cumplimiento_epp.feature: traducen el texto y leen el Excel."""

import functools
import sqlite3
from datetime import UTC, date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pytest_bdd import given, parsers, scenarios, then, when

from src.analysis.reporte_cumplimiento_epp import GeneradorReporteCumplimiento, PedidoReporte
from src.ppe.domain.elementos import ElementoEPP
from src.ppe.domain.infracciones import Incumplimiento, InfraccionFinalizada, InfraccionIniciada
from src.storage.infracciones_epp import RepositorioInfraccionesSQLite, crear_tabla_infracciones
from tests.dobles_ppe import RelojFalso

pytestmark = pytest.mark.bdd

scenarios("../reporte_cumplimiento_epp.feature")

SUCURSAL = timezone(timedelta(hours=-3))
DIA = date(2026, 10, 22)
DIA_SIGUIENTE = datetime(2026, 10, 23, 12, 0, tzinfo=UTC)
FILA_DE_DATOS = 3
ZONA = "Boca de pozo"


class Sucursal:
    """La base de la sucursal en tmp_path y el reporte generado."""

    def __init__(self, carpeta: Path) -> None:
        self.conectar = functools.partial(sqlite3.connect, carpeta / "tracking.db")
        crear_tabla_infracciones(self.conectar)
        with self.conectar() as conexion:
            conexion.execute("CREATE TABLE employees (employee_name TEXT, shift TEXT)")
        self.reporte = carpeta / "cumplimiento.xlsx"
        self.turnos: dict[str, str] = {}

    def registrar(self, nombre: str, elemento: str, hora: str, minutos: int) -> None:
        inicio = datetime.combine(DIA, datetime.strptime(hora, "%H:%M").time(), SUCURSAL)
        incumplimiento = Incumplimiento(1, nombre, ElementoEPP(elemento), ZONA)
        repositorio = RepositorioInfraccionesSQLite(self.conectar)
        repositorio.guardar(InfraccionIniciada(incumplimiento, inicio, "epp", None))
        fin = inicio + timedelta(minutes=minutos)
        repositorio.guardar(InfraccionFinalizada(incumplimiento, inicio, "epp", fin))

    def registrar_presencia(self, nombre: str, hora: str, minutos: int) -> None:
        """Un registro del seguimiento por minuto, como lo deja la cámara con la persona en zona."""
        inicio = datetime.combine(DIA, datetime.strptime(hora, "%H:%M").time())
        horas = [(inicio + timedelta(minutes=minuto)).isoformat() for minuto in range(minutos + 1)]
        with self.conectar() as conexion:
            conexion.execute(
                "CREATE TABLE IF NOT EXISTS tracking "
                "(employee_name TEXT, timestamp TEXT, zone TEXT, inside_zone INTEGER)"
            )
            conexion.executemany(
                "INSERT INTO tracking VALUES (?, ?, ?, 1)", [(nombre, h, ZONA) for h in horas]
            )

    def filas(self) -> list[tuple[object, ...]]:
        hoja = load_workbook(self.reporte)["Cumplimiento EPP"]
        return list(hoja.iter_rows(min_row=FILA_DE_DATOS, values_only=True))


@pytest.fixture
def sucursal(tmp_path: Path) -> Sucursal:
    return Sucursal(tmp_path)


@given(parsers.parse('que "{nombre}" trabaja en el turno "{turno}"'))
def empleado_con_turno(sucursal: Sucursal, nombre: str, turno: str) -> None:
    with sucursal.conectar() as conexion:
        conexion.execute("INSERT INTO employees VALUES (?, ?)", (nombre, turno))
    sucursal.turnos[nombre] = turno


@given(parsers.parse('que "{nombre}" estuvo sin {elemento} {minutos:d} minutos a las {hora}'))
def infraccion(sucursal: Sucursal, nombre: str, elemento: str, minutos: int, hora: str) -> None:
    sucursal.registrar(nombre, elemento, hora, minutos)


@given(parsers.parse('que "{nombre}" estuvo {minutos:d} minutos en su zona desde las {hora}'))
def presencia(sucursal: Sucursal, nombre: str, minutos: int, hora: str) -> None:
    sucursal.registrar_presencia(nombre, hora, minutos)


@given("que no hubo infracciones en el día")
def sin_infracciones() -> None:
    """La base empieza vacía."""


@when("genero el reporte de cumplimiento del día")
def generar(sucursal: Sucursal) -> None:
    generador = GeneradorReporteCumplimiento(sucursal.conectar, RelojFalso(DIA_SIGUIENTE), SUCURSAL)
    generador.generar(PedidoReporte(DIA, DIA, sucursal.reporte))


@then(
    parsers.parse(
        'el reporte muestra para "{nombre}", en su turno, {cantidad:d} infracciones '
        "sin {elemento} que suman {minutos:d} minutos"
    )
)
def fila_del_reporte(
    sucursal: Sucursal, nombre: str, cantidad: int, elemento: str, minutos: int
) -> None:
    filas = [fila[:6] for fila in sucursal.filas()]
    turno = sucursal.turnos[nombre]

    assert filas == [(turno, nombre, ZONA, elemento, cantidad, float(minutos))]


@then("el reporte indica que no hubo infracciones en el período")
def reporte_vacio(sucursal: Sucursal) -> None:
    assert sucursal.filas() == [("Sin infracciones en el período",) + (None,) * 6]


def _cumplimiento(sucursal: Sucursal, nombre: str, elemento: str) -> object:
    (fila,) = (f for f in sucursal.filas() if f[1] == nombre and f[3] == elemento)
    return fila[6]


@then(parsers.parse('el cumplimiento de {elemento} de "{nombre}" es del {porcentaje:d} %'))
def cumplimiento(sucursal: Sucursal, elemento: str, nombre: str, porcentaje: int) -> None:
    assert _cumplimiento(sucursal, nombre, elemento) == float(porcentaje)


@then(parsers.parse('el cumplimiento de {elemento} de "{nombre}" figura sin datos'))
def cumplimiento_sin_datos(sucursal: Sucursal, elemento: str, nombre: str) -> None:
    assert _cumplimiento(sucursal, nombre, elemento) == "sin datos"
