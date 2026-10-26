"""Reporte Excel de cumplimiento de EPP (SPEC/9 §8.4; CLAUDE.md §16).

Se verifica el contenido de las celdas, no el tamaño del archivo. El período son días completos
en el huso horario de la sucursal; el reporte se genera también en segundo plano.
"""

import functools
import sqlite3
import threading
from datetime import UTC, date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.analysis.reporte_cumplimiento_epp import (
    GeneradorReporteCumplimiento,
    PedidoReporte,
    periodo,
)
from src.ppe.domain.elementos import ElementoEPP
from src.ppe.domain.infracciones import Incumplimiento, InfraccionFinalizada, InfraccionIniciada
from src.storage.infracciones_epp import RepositorioInfraccionesSQLite, crear_tabla_infracciones
from tests.dobles_ppe import RelojFalso

Conectar = functools.partial[sqlite3.Connection]
ARGENTINA = timezone(timedelta(hours=-3))
DIA = date(2026, 10, 22)
# Medianoche del 23/10 en Argentina: el reporte del 22/10 ya cerró.
DESPUES_DEL_DIA = datetime(2026, 10, 23, 3, 0, tzinfo=UTC)
ESPERA_DEL_HILO = 10.0
TITULO = (
    "Reporte de cumplimiento de EPP",
    "del 2026-10-22 al 2026-10-22",
    *(None,) * 5,
)
ENCABEZADOS = (
    "Turno",
    "Empleado",
    "Zona",
    "Elemento",
    "Infracciones",
    "Duración total (min)",
    "% del tiempo en cumplimiento",
)


@pytest.fixture
def conectar(tmp_path: Path) -> Conectar:
    conectar = functools.partial(sqlite3.connect, tmp_path / "tracking.db")
    crear_tabla_infracciones(conectar)
    with conectar() as conexion:
        conexion.execute("CREATE TABLE employees (employee_name TEXT PRIMARY KEY, shift TEXT)")
        conexion.execute("INSERT INTO employees VALUES ('Ana', 'Matutino')")
    return conectar


def _registrar(
    conectar: Conectar, nombre: str, inicio: datetime, minutos: int, elemento: ElementoEPP
) -> None:
    incumplimiento = Incumplimiento(1, nombre, elemento, "Boca de pozo")
    repositorio = RepositorioInfraccionesSQLite(conectar)
    repositorio.guardar(InfraccionIniciada(incumplimiento, inicio, "epp-0.1.0", None))
    fin = inicio + timedelta(minutes=minutos)
    repositorio.guardar(InfraccionFinalizada(incumplimiento, inicio, "epp-0.1.0", fin))


def _local(hora: int, minuto: int = 0, dia: int = 22) -> datetime:
    return datetime(2026, 10, dia, hora, minuto, tzinfo=ARGENTINA)


def _generador(
    conectar: Conectar, ahora: datetime = DESPUES_DEL_DIA
) -> GeneradorReporteCumplimiento:
    return GeneradorReporteCumplimiento(conectar, RelojFalso(ahora), ARGENTINA)


def _celdas(ruta: Path) -> list[tuple[object, ...]]:
    libro = load_workbook(ruta)
    hoja = libro["Cumplimiento EPP"]
    return [tuple(fila) for fila in hoja.iter_rows(values_only=True)]


def test_periodo_son_dias_completos_en_el_huso_de_la_sucursal() -> None:
    pedido = PedidoReporte(DIA, DIA, Path("r.xlsx"))

    assert periodo(pedido, ARGENTINA) == (_local(0), _local(0, dia=23))


def test_periodo_con_las_fechas_invertidas_falla() -> None:
    pedido = PedidoReporte(DIA, DIA - timedelta(days=1), Path("r.xlsx"))

    with pytest.raises(
        ValueError, match=r"^La fecha desde \(2026-10-22\) es posterior a hasta \(2026-10-21\)$"
    ):
        periodo(pedido, ARGENTINA)


def test_reporte_escribe_una_fila_por_turno_empleado_zona_y_elemento(
    conectar: Conectar, tmp_path: Path
) -> None:
    _registrar(conectar, "Ana", _local(8), 10, ElementoEPP.CASCO)
    _registrar(conectar, "Ana", _local(9), 5, ElementoEPP.CASCO)
    _registrar(conectar, "Luis", _local(23, 50), 20, ElementoEPP.CHALECO)
    _registrar(conectar, "Ana", _local(0, 30, dia=23), 5, ElementoEPP.CASCO)
    ruta = tmp_path / "cumplimiento.xlsx"

    filas = _generador(conectar).generar(PedidoReporte(DIA, DIA, ruta))

    # Luis empezó a las 23:50 y siguió 20 minutos: en el reporte del 22 cuentan 10. Lo de Ana
    # a las 00:30 del 23 es del día siguiente.
    assert (filas, _celdas(ruta)) == (
        2,
        [
            TITULO,
            ENCABEZADOS,
            ("Matutino", "Ana", "Boca de pozo", "casco", 2, 15.0, "sin datos"),
            ("sin turno", "Luis", "Boca de pozo", "chaleco", 1, 10.0, "sin datos"),
        ],
    )


def test_reporte_calcula_el_porcentaje_con_el_tiempo_en_la_zona(
    conectar: Conectar, tmp_path: Path
) -> None:
    # Ana estuvo 20 minutos en la zona (un registro por minuto) y 5 sin casco: 75 % cumplido.
    with conectar() as conexion:
        conexion.execute(
            "CREATE TABLE tracking (employee_name TEXT, timestamp TEXT, zone TEXT, "
            "inside_zone INTEGER)"
        )
        conexion.executemany(
            "INSERT INTO tracking VALUES ('Ana', ?, 'Boca de pozo', 1)",
            [(f"2026-10-22T08:{minuto:02d}:00",) for minuto in range(21)],
        )
    _registrar(conectar, "Ana", _local(8), 5, ElementoEPP.CASCO)
    ruta = tmp_path / "con_tiempo.xlsx"

    _generador(conectar).generar(PedidoReporte(DIA, DIA, ruta))

    assert _celdas(ruta)[2] == ("Matutino", "Ana", "Boca de pozo", "casco", 1, 5.0, 75.0)


def test_reporte_de_hoy_cuenta_las_infracciones_abiertas_hasta_ahora(
    conectar: Conectar, tmp_path: Path
) -> None:
    incumplimiento = Incumplimiento(1, "Ana", ElementoEPP.CASCO, "Boca de pozo")
    evento = InfraccionIniciada(incumplimiento, _local(8), "epp-0.1.0", None)
    RepositorioInfraccionesSQLite(conectar).guardar(evento)
    ruta = tmp_path / "hoy.xlsx"

    ahora = _local(8, 30) + timedelta(seconds=20)
    _generador(conectar, ahora=ahora).generar(PedidoReporte(DIA, DIA, ruta))

    assert _celdas(ruta)[2][4:6] == (1, 30.3)


def test_reporte_sin_infracciones_lo_dice(conectar: Conectar, tmp_path: Path) -> None:
    ruta = tmp_path / "vacio.xlsx"

    filas = _generador(conectar).generar(PedidoReporte(DIA, DIA, ruta))

    assert filas == 0
    assert _celdas(ruta)[2][0] == "Sin infracciones en el período"


def test_reporte_en_segundo_plano_avisa_al_terminar(conectar: Conectar, tmp_path: Path) -> None:
    _registrar(conectar, "Ana", _local(8), 10, ElementoEPP.CASCO)
    ruta = tmp_path / "async.xlsx"
    resultados: list[tuple[Path, int]] = []
    errores: list[Exception] = []

    hilo = _generador(conectar).generar_en_segundo_plano(
        PedidoReporte(DIA, DIA, ruta),
        lambda ruta, filas: resultados.append((ruta, filas)),
        errores.append,
    )
    hilo.join(ESPERA_DEL_HILO)

    assert (resultados, errores) == ([(ruta, 1)], [])
    assert (hilo.daemon, hilo.name) == (True, "reporte-cumplimiento-epp")


def test_reporte_en_segundo_plano_avisa_si_falla(conectar: Conectar, tmp_path: Path) -> None:
    errores: list[Exception] = []
    pedido = PedidoReporte(DIA, DIA, tmp_path / "no-existe" / "reporte.xlsx")

    hilo = _generador(conectar).generar_en_segundo_plano(
        pedido, lambda ruta, filas: pytest.fail("no debía terminar bien"), errores.append
    )
    hilo.join(ESPERA_DEL_HILO)

    assert [type(error) for error in errores] == [FileNotFoundError]


def test_reporte_no_bloquea_al_que_lo_pide(conectar: Conectar, tmp_path: Path) -> None:
    liberar = threading.Event()
    errores: list[Exception] = []
    generador = _generador(conectar)
    pedido = PedidoReporte(DIA, DIA, tmp_path / "r.xlsx")

    hilo = generador.generar_en_segundo_plano(
        pedido, lambda ruta, filas: liberar.wait(ESPERA_DEL_HILO), errores.append
    )

    assert hilo.is_alive() is True
    liberar.set()
    hilo.join(ESPERA_DEL_HILO)
    assert errores == []
