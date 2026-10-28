"""Sucursal simulada de punta a punta para el recorrido del comprador (SPEC/9 §16).

Encadena las mismas piezas que arma la aplicación, sin cámara ni modelo reales (CLAUDE.md §16):
el polígono se guarda como lo hace el editor de zonas, el EPP exigido se marca con el
presentador de la vista "Zonas", las zonas de cada persona salen de ZoneChecker sobre ese mismo
archivo (como en TrackingPipeline) y el caso de uso se arma con `crear_evaluador_epp` y los
parámetros que decidió el negocio. La base, las evidencias y el Excel viven en `tmp_path`; el
webhook es un servidor falso en 127.0.0.1 y el reloj es falso.
"""

import sqlite3
from collections.abc import Callable, Sequence
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import NamedTuple

import numpy as np
from openpyxl import load_workbook

from src.analysis.reporte_cumplimiento_epp import GeneradorReporteCumplimiento, PedidoReporte
from src.composicion_avisos import preparar_avisos
from src.composicion_epp import (
    PARAMETROS_DECIDIDOS,
    DestinoEPP,
    PublicadorEnMemoria,
    crear_evaluador_epp,
)
from src.gui.presentador_alertas import FilaAlerta, PreferenciasPanel, PresentadorAlertas
from src.gui.presentador_zonas import PresentadorZonas
from src.inference.types import Caja, ClaseObjeto, Deteccion
from src.ppe.domain.elementos import ElementoEPP
from src.ppe.domain.infracciones import EventoInfraccion
from src.ppe.domain.reglas import ReglasEPP
from src.storage.alertas_panel import RepositorioAlertasSQLite
from src.tracking.integracion_epp import IntegracionEPP, personas_en_escena
from src.zones.edicion_zonas import guardar_poligono
from src.zones.formato_zonas import leer_zonas
from src.zones.zone_checker import ZoneChecker
from tests.dobles_ppe import RelojFalso
from tests.escenas_ppe import CABEZA, EPP_COLOCADO, PERSONA, Rectangulo

INICIO_DEL_TURNO = datetime(2026, 10, 26, 8, 0, tzinfo=UTC)
UN_SEGUNDO = timedelta(seconds=1)
VERSION_DEL_MODELO = "epp-1.0.0"
SECRETO_DEL_WEBHOOK = "secreto-de-prueba"  # pragma: allowlist secret
SUPERVISOR = "supervisor"
LIMITE_DEL_PANEL = 200
ESPERA_DEL_DESPACHADOR = timedelta(seconds=10)
CONFIANZA = 0.9
# D3: la vista "Zonas" ofrece casco y chaleco en la primera versión (src/gui/views.py).
OFRECIDOS = (ElementoEPP.CASCO, ElementoEPP.CHALECO)
# Polígono de 400 x 400 píxeles en la esquina del cuadro; fuera de él, x >= 500.
BOCA_DE_POZO = "Boca de pozo"
POLIGONO_BOCA_DE_POZO = ((0.0, 0.0), (400.0, 0.0), (400.0, 400.0), (0.0, 400.0))
DENTRO_DE_LA_ZONA = 100.0
FUERA_DE_TODA_ZONA = 600.0


class Operario(NamedTuple):
    """Persona frente a la cámara: el seguimiento le da un track; el rostro, quizá, un nombre."""

    nombre: str | None
    track_id: int
    x: float
    lleva: frozenset[ElementoEPP]

    def detecciones(self) -> list[Deteccion]:
        """Lo que el detector vería en el cuadro: persona, cabeza y el EPP que lleva puesto."""
        cajas = [(ClaseObjeto.PERSONA, PERSONA), (ClaseObjeto.CABEZA, CABEZA)]
        cajas += [
            (ClaseObjeto(elemento.value), rectangulo)
            for elemento in ElementoEPP
            if elemento in self.lleva
            for rectangulo in EPP_COLOCADO[elemento]
        ]
        return [Deteccion(clase, CONFIANZA, self._caja(r)) for clase, r in cajas]

    def _caja(self, rectangulo: Rectangulo) -> Caja:
        x1, y1, x2, y2 = rectangulo
        return Caja(self.x + x1, y1, self.x + x2, y2)


def operario(
    nombre: str | None, *lleva: ElementoEPP, x: float = DENTRO_DE_LA_ZONA, track_id: int = 1
) -> Operario:
    return Operario(nombre, track_id, x, frozenset(lleva))


class SucursalSimulada:
    """Una sucursal con su base, sus zonas, una cámara y el panel del responsable."""

    def __init__(self, carpeta: Path) -> None:
        ruta_base = carpeta / "db" / "local_tracking.db"
        ruta_base.parent.mkdir(parents=True)
        self.conectar: Callable[[], sqlite3.Connection] = lambda: sqlite3.connect(ruta_base)
        self.ruta_zonas = carpeta / "zonas" / "zonas_config.json"
        self.carpeta_snapshots = carpeta / "snapshots"
        self.carpeta_export = carpeta / "export"
        self.carpeta_export.mkdir()
        self.reloj = RelojFalso(INICIO_DEL_TURNO)
        self.publicados: list[EventoInfraccion] = []
        self._camara: IntegracionEPP | None = None
        self._zonas: ZoneChecker | None = None

    # --- Configuración (vista "Zonas" y editor de polígonos) -----------------------------------

    def dibujar_zona(self, nombre: str, puntos: Sequence[tuple[float, float]]) -> None:
        guardar_poligono(self.ruta_zonas, nombre, puntos)

    def exigir(self, zona: str, *elementos: ElementoEPP) -> str | None:
        """Marca las casillas de EPP de la zona y guarda; devuelve el mensaje de la vista."""
        presentador = PresentadorZonas(self.ruta_zonas, OFRECIDOS)
        presentador.cargar()
        for elemento in elementos:
            presentador.alternar(zona, elemento)
        presentador.guardar()
        return presentador.mensaje

    # --- Monitoreo (una cámara con el módulo de EPP) -------------------------------------------

    def conectar_camara(self) -> None:
        """Como crear_integracion_epp_de_la_app: las zonas se leen al conectar la cámara."""
        publicador = PublicadorEnMemoria()
        publicador.suscribir(self.publicados.append)
        destino = DestinoEPP(self.conectar, self.carpeta_snapshots, publicador)
        reglas = ReglasEPP(frozenset(), leer_zonas(self.ruta_zonas))
        evaluador = crear_evaluador_epp(destino, reglas, VERSION_DEL_MODELO, self.reloj)
        self._camara = IntegracionEPP(evaluador, PARAMETROS_DECIDIDOS.cada_n_frames)
        self._zonas = ZoneChecker(str(self.ruta_zonas))

    def filmar(self, operarios: Sequence[Operario], duracion: timedelta) -> None:
        """Un cuadro por segundo durante `duracion`, más el cuadro final."""
        for _ in range(int(duracion / UN_SEGUNDO) + 1):
            self._cuadro(operarios)
            self.reloj.avanzar(UN_SEGUNDO)

    def _cuadro(self, operarios: Sequence[Operario]) -> None:
        assert self._camara is not None, "Primero hay que conectar la cámara"
        imagen = np.zeros((400, 800, 3), dtype=np.uint8)
        detecciones = [d for o in operarios for d in o.detecciones()]
        self._camara.evaluar(imagen, personas_en_escena(self._seguimiento(operarios)), detecciones)

    def _seguimiento(self, operarios: Sequence[Operario]) -> list[dict[str, object]]:
        """Lo que TrackingPipeline.process entrega por persona (zonas por el centro de la caja)."""
        assert self._zonas is not None
        datos = []
        for o in operarios:
            caja = next(d.caja for d in o.detecciones() if d.clase is ClaseObjeto.PERSONA)
            zonas = self._zonas.check((caja.x1 + caja.x2) / 2, (caja.y1 + caja.y2) / 2)
            datos.append(
                {
                    "track_id": o.track_id,
                    "bbox": (caja.x1, caja.y1, caja.x2, caja.y2),
                    "zones": frozenset(nombre for nombre, dentro in zonas.items() if dentro),
                    "name": o.nombre or "Unknown",
                }
            )
        return datos

    # --- Avisos (apertura de la sucursal con un webhook configurado) ---------------------------

    def abrir_con_webhook(self, url: str) -> None:
        """Lo que hace la app al abrir la sucursal: purga, arma el despachador y despacha."""
        entorno = {
            "EPP_ALERTAS_WEBHOOK_URL": url,
            "EPP_ALERTAS_WEBHOOK_SECRETO": SECRETO_DEL_WEBHOOK,
            "EPP_ALERTAS_WEBHOOK_TIMEOUT_S": "5",
        }
        despachador = preparar_avisos(entorno, self.conectar, self.carpeta_snapshots, self.reloj)
        assert despachador is not None, "Con un webhook configurado debe haber despachador"
        despachador.iniciar()
        despachador.detener(ESPERA_DEL_DESPACHADOR)

    def abrir_sin_canal(self) -> None:
        """Abrir la sucursal sin canal de aviso: solo prepara la base y purga lo vencido."""
        despachador = preparar_avisos({}, self.conectar, self.carpeta_snapshots, self.reloj)
        assert despachador is None, "Sin canal configurado no debe haber despachador"

    # --- Consultas del responsable --------------------------------------------------------------

    def panel(self) -> PresentadorAlertas:
        preferencias = PreferenciasPanel(SUPERVISOR, UTC, LIMITE_DEL_PANEL)
        presentador = PresentadorAlertas(
            RepositorioAlertasSQLite(self.conectar), self.reloj, preferencias
        )
        presentador.actualizar()
        return presentador

    def alertas(self) -> list[FilaAlerta]:
        return self.panel().filas

    def reporte_del_dia(self, dia: date) -> list[tuple[object, ...]]:
        """Filas de datos del Excel de cumplimiento (sin título ni encabezados)."""
        ruta = self.carpeta_export / f"Cumplimiento_EPP_{dia}_a_{dia}.xlsx"
        generador = GeneradorReporteCumplimiento(self.conectar, self.reloj, UTC)
        generador.generar(PedidoReporte(dia, dia, ruta))
        hoja = load_workbook(ruta)["Cumplimiento EPP"]
        return [tuple(fila) for fila in hoja.iter_rows(min_row=3, values_only=True)]

    def evidencias(self) -> list[Path]:
        return sorted((self.carpeta_snapshots / "epp").rglob("*.jpg"))

    def registros(self) -> list[tuple[object, ...]]:
        consulta = (
            "SELECT employee_name, zona, elemento, estado_notificacion, evidencia_path "
            "FROM ppe_infracciones ORDER BY id"
        )
        conexion = self.conectar()
        filas = conexion.execute(consulta).fetchall()
        conexion.close()
        return filas
